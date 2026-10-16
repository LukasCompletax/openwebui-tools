"""
title: Beratungsvertrag als Word-Datei (Vorlage befüllen)
author: completax
version: 0.1.0
description: Befüllt die Word-Vorlage des Mandatsvertrags (Beratungsvertrag, Angebot über steuerliche Beratungsleistungen) mit Mandantendaten, Pauschalen und Bedingungen und liefert die fertige DOCX zum Download. Das Modell liefert nur die Werte, das Tool setzt sie deterministisch in die Vorlage ein.
requirements: python-docx
"""

from __future__ import annotations

import asyncio
import datetime as dt
import io
import json
import logging
import os
import re
import tempfile
import time
import uuid
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from typing import Any, Optional

from pydantic import BaseModel, Field

log = logging.getLogger(__name__)

DEFAULT_TEMPLATE_PATH = "/app/backend/data/templates/beratungsvertrag/vorlage.docx"
DOCX_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"

_PLACEHOLDER = re.compile(r"\{\{([A-Za-z0-9_]+)\}\}")
_BLOCK_START = "{{#%s}}"
_BLOCK_END = "{{/%s}}"

# Platzhalter der Vorlage. Die Vorlage selbst liegt nicht im Repository.
PLACEHOLDERS = (
    "Mandant_Name",
    "Mandant_Anschrift",
    "Mandant_Email",
    "Mandant_Kurzname",
    "Jahr",
    "Vertragsdatum",
    "Gespraechsdatum",
    "Erste_Anpassung_Datum",
    "Anzahlung_Prozent",
    "Buchungen_Anzahl",
    "Pauschale_Buchfuehrung_Monat",
    "Lohn_Einrichtung",
    "Pauschale_Jahresabschluss",
    "Erstes_Geschaeftsjahr",
    "Pauschale_Steuererklaerung",
)
BLOCKS = ("Anzahlung", "Verbraucher")


# ---------------------------------------------------------------------------
# Werte: Datum, Beträge, Prozent
# ---------------------------------------------------------------------------


def parse_date(value: str) -> dt.date:
    """TT.MM.JJJJ oder JJJJ-MM-TT. Alles andere ist ein Fehler."""
    text = (value or "").strip()
    for fmt in ("%d.%m.%Y", "%Y-%m-%d"):
        try:
            return dt.datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    raise ValueError(f"Datum '{text}' nicht lesbar. Erwartet TT.MM.JJJJ.")


def format_date(value: dt.date) -> str:
    return value.strftime("%d.%m.%Y")


def erste_anpassung(vertragsdatum: dt.date) -> dt.date:
    """Vertragsdatum plus ein Jahr, danach der erste 1. Januar, der folgt oder zusammenfällt."""
    year = vertragsdatum.year + 1
    if vertragsdatum.month == 1 and vertragsdatum.day == 1:
        return dt.date(year, 1, 1)
    return dt.date(year + 1, 1, 1)


def parse_amount(value: str) -> Decimal:
    """Betrag in EUR lesen: '450', '450,00', '1.250,00', '1250.5', 'EUR 450'. Ergebnis exakt als Decimal."""
    text = (value or "").strip()
    text = re.sub(r"(?i)\b(eur|euro)\b|€", "", text).replace(" ", "").replace(" ", "")
    if not text:
        raise ValueError("Betrag fehlt.")
    if "," in text and "." in text:
        if text.rfind(",") > text.rfind("."):
            text = text.replace(".", "").replace(",", ".")  # 1.250,00
        else:
            text = text.replace(",", "")  # 1,250.00
    elif "," in text:
        text = text.replace(",", ".")  # 450,00
    elif text.count(".") == 1 and re.fullmatch(r"-?\d{1,3}\.\d{3}", text):
        text = text.replace(".", "")  # 1.250 (deutsche Tausender)
    elif text.count(".") > 1:
        text = text.replace(".", "")  # 1.250.000
    try:
        amount = Decimal(text)
    except InvalidOperation as exc:
        raise ValueError(f"Betrag '{value}' nicht lesbar.") from exc
    if amount < 0:
        raise ValueError(f"Betrag '{value}' darf nicht negativ sein.")
    return amount.quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)


def format_amount(amount: Decimal) -> str:
    """Decimal → '1.250,00' (ohne Währungszeichen, so wie es die Vorlage erwartet)."""
    text = f"{amount:,.2f}"
    return text.replace(",", "X").replace(".", ",").replace("X", ".")


def parse_percent(value: str) -> Optional[Decimal]:
    """Prozentsatz der Anzahlung. Leer, '0' oder 'nein' bedeutet keine Anzahlung."""
    text = (value or "").strip().lower().replace("%", "").strip()
    if text in ("", "0", "nein", "keine", "false", "none"):
        return None
    try:
        percent = Decimal(text.replace(",", "."))
    except InvalidOperation as exc:
        raise ValueError(f"Anzahlung '{value}' nicht lesbar. Erwartet eine Zahl in Prozent.") from exc
    if not 0 < percent <= 100:
        raise ValueError(f"Anzahlung '{value}' liegt nicht zwischen 0 und 100 Prozent.")
    return percent


def format_percent(percent: Decimal) -> str:
    if percent == percent.to_integral_value():
        return str(int(percent))
    return str(percent.normalize()).replace(".", ",")


def parse_year(value: str, field: str) -> str:
    text = (value or "").strip()
    if not re.fullmatch(r"(19|20)\d{2}", text):
        raise ValueError(f"{field}: '{text}' ist kein Jahr (JJJJ).")
    return text


def safe_file_name(name: str, fallback: str = "Mandatsvertrag") -> str:
    base = os.path.basename((name or "").strip())
    base = re.sub(r"(?i)\.docx$", "", base)
    base = re.sub(r"[^\w\-. ]+", "_", base, flags=re.UNICODE).strip(" ._")[:80] or fallback
    return f"{base}.docx"


# ---------------------------------------------------------------------------
# DOCX-Vorlage befüllen (run-sicher, Blöcke über Absatzgrenzen, Zeilen streichen)
# ---------------------------------------------------------------------------


def _runs_text(paragraph: Any) -> str:
    return "".join(run.text for run in paragraph.runs)


def _replace_span(paragraph: Any, start: int, end: int, replacement: str) -> None:
    """Zeichen [start, end) im Absatz ersetzen, Formatierung der Runs bleibt erhalten."""
    pos = 0
    first = True
    for run in paragraph.runs:
        run_start, run_end = pos, pos + len(run.text)
        pos = run_end
        if run_end <= start or run_start >= end:
            continue
        lo = max(start, run_start) - run_start
        hi = min(end, run_end) - run_start
        text = run.text
        run.text = text[:lo] + (replacement if first else "") + text[hi:]
        first = False


def _delete_paragraph(paragraph: Any) -> None:
    element = paragraph._element
    element.getparent().remove(element)


def _containers(document: Any) -> list[list[Any]]:
    """Absatzlisten in Dokumentreihenfolge: Textkörper, Tabellenzellen, Kopf- und Fußzeilen."""
    groups: list[list[Any]] = [list(document.paragraphs)]
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                groups.append(list(cell.paragraphs))
    for section in document.sections:
        for part in (section.header, section.footer):
            groups.append(list(part.paragraphs))
            for table in part.tables:
                for row in table.rows:
                    for cell in row.cells:
                        groups.append(list(cell.paragraphs))
    return groups


def _apply_block(paragraphs: list[Any], name: str, keep: bool) -> int:
    """Block {{#name}} … {{/name}} behalten (Marker entfernen) oder samt Inhalt löschen."""
    start_marker, end_marker = _BLOCK_START % name, _BLOCK_END % name
    start = end = None
    for index, paragraph in enumerate(paragraphs):
        text = _runs_text(paragraph)
        if start is None and start_marker in text:
            start = (index, text.index(start_marker))
        if start is not None and end_marker in text:
            offset = text.index(end_marker)
            if index > start[0] or offset >= start[1] + len(start_marker):
                end = (index, offset)
                break
    if start is None and end is None:
        return 0
    if start is None or end is None:
        raise ValueError(f"Block '{name}' ist in der Vorlage unvollständig (Anfang oder Ende fehlt).")

    (s_idx, s_off), (e_idx, e_off) = start, end
    if keep:
        _replace_span(paragraphs[e_idx], e_off, e_off + len(end_marker), "")
        _replace_span(paragraphs[s_idx], s_off, s_off + len(start_marker), "")
    elif s_idx == e_idx:
        _replace_span(paragraphs[s_idx], s_off, e_off + len(end_marker), "")
    else:
        _replace_span(paragraphs[e_idx], 0, e_off + len(end_marker), "")
        _replace_span(paragraphs[s_idx], s_off, len(_runs_text(paragraphs[s_idx])), "")
        for paragraph in paragraphs[s_idx + 1 : e_idx]:
            _delete_paragraph(paragraph)
    removed = 0
    for index in (e_idx, s_idx):
        if not _runs_text(paragraphs[index]).strip():
            _delete_paragraph(paragraphs[index])
            removed += 1
    return removed


def fill_docx(
    template_path: str,
    values: dict[str, str],
    blocks: dict[str, bool],
    drop_paragraphs_with: list[str],
    out_path: str,
) -> dict:
    """Vorlage befüllen. Reihenfolge: Blöcke, Zeilen streichen, Platzhalter ersetzen, Restkontrolle."""
    import docx

    document = docx.Document(template_path)
    stats = {"replaced": 0, "paragraphs_removed": 0, "blocks": {}, "remaining": []}

    for name, keep in blocks.items():
        for paragraphs in _containers(document):
            stats["paragraphs_removed"] += _apply_block(paragraphs, name, keep)
        stats["blocks"][name] = "behalten" if keep else "entfernt"

    markers = [f"{{{{{name}}}}}" for name in drop_paragraphs_with]
    for paragraphs in _containers(document):
        for paragraph in paragraphs:
            text = _runs_text(paragraph)
            if any(marker in text for marker in markers):
                _delete_paragraph(paragraph)
                stats["paragraphs_removed"] += 1

    for paragraphs in _containers(document):
        for paragraph in paragraphs:
            text = _runs_text(paragraph)
            for match in reversed(list(_PLACEHOLDER.finditer(text))):
                name = match.group(1)
                if name in values:
                    _replace_span(paragraph, match.start(), match.end(), values[name])
                    stats["replaced"] += 1

    remaining: set[str] = set()
    for paragraphs in _containers(document):
        for paragraph in paragraphs:
            remaining.update(re.findall(r"\{\{[^}]*\}\}", paragraph.text))
    stats["remaining"] = sorted(remaining)
    document.save(out_path)
    return stats


# ---------------------------------------------------------------------------
# Fachlogik des Mandatsvertrags: Eingaben → Platzhalterwerte
# ---------------------------------------------------------------------------


def build_contract(
    mandant_name: str,
    mandant_anschrift: str,
    gespraechsdatum: str,
    mandant_email: str = "",
    mandant_kurzname: str = "",
    jahr: str = "",
    vertragsdatum: str = "",
    verbraucher: bool = False,
    anzahlung_prozent: str = "",
    buchungen_anzahl: int = 0,
    pauschale_buchfuehrung_monat: str = "",
    lohn_einrichtung: str = "",
    pauschale_jahresabschluss: str = "",
    erstes_geschaeftsjahr: str = "",
    pauschale_steuererklaerung: str = "",
    today: Optional[dt.date] = None,
) -> dict:
    """Prüft die Eingaben und liefert values, blocks, drops, summary. Wirft ValueError mit allen Mängeln."""
    errors: list[str] = []
    values: dict[str, str] = {}

    name = (mandant_name or "").strip()
    anschrift = " ".join((mandant_anschrift or "").split())
    if not name:
        errors.append("mandant_name fehlt.")
    if not anschrift:
        errors.append("mandant_anschrift fehlt.")

    vertrag = today or dt.date.today()
    if (vertragsdatum or "").strip():
        try:
            vertrag = parse_date(vertragsdatum)
        except ValueError as exc:
            errors.append(f"vertragsdatum: {exc}")
    try:
        gespraech = parse_date(gespraechsdatum)
    except ValueError as exc:
        gespraech = None
        errors.append(f"gespraechsdatum: {exc}")

    jahr_text = (jahr or "").strip() or str(vertrag.year)
    try:
        jahr_text = parse_year(jahr_text, "jahr")
    except ValueError as exc:
        errors.append(str(exc))

    try:
        anzahlung = parse_percent(anzahlung_prozent)
    except ValueError as exc:
        anzahlung = None
        errors.append(str(exc))

    pauschalen: dict[str, Optional[Decimal]] = {}
    for key, raw in (
        ("Pauschale_Buchfuehrung_Monat", pauschale_buchfuehrung_monat),
        ("Lohn_Einrichtung", lohn_einrichtung),
        ("Pauschale_Jahresabschluss", pauschale_jahresabschluss),
        ("Pauschale_Steuererklaerung", pauschale_steuererklaerung),
    ):
        text = (raw or "").strip()
        if not text or text.lower() in ("0", "nein", "keine", "none"):
            pauschalen[key] = None
            continue
        try:
            pauschalen[key] = parse_amount(text)
        except ValueError as exc:
            errors.append(f"{key.lower()}: {exc}")
            pauschalen[key] = None

    if pauschalen["Pauschale_Buchfuehrung_Monat"] is not None and int(buchungen_anzahl or 0) <= 0:
        errors.append("buchungen_anzahl fehlt: Für die Finanzbuchführungspauschale ist die Zahl der Buchungen je Monat nötig.")
    geschaeftsjahr = ""
    if pauschalen["Pauschale_Jahresabschluss"] is not None:
        try:
            geschaeftsjahr = parse_year(erstes_geschaeftsjahr, "erstes_geschaeftsjahr")
        except ValueError:
            errors.append("erstes_geschaeftsjahr fehlt oder ist kein Jahr: Für die Jahresabschlusspauschale ist das erste Geschäftsjahr nötig (JJJJ).")

    if errors:
        raise ValueError(" ".join(errors))

    anpassung = erste_anpassung(vertrag)
    kurzname = (mandant_kurzname or "").strip() or name
    email = (mandant_email or "").strip()

    values.update(
        {
            "Mandant_Name": name,
            "Mandant_Anschrift": anschrift,
            "Mandant_Email": email,
            "Mandant_Kurzname": kurzname,
            "Jahr": jahr_text,
            "Vertragsdatum": format_date(vertrag),
            "Gespraechsdatum": format_date(gespraech),  # type: ignore[arg-type]
            "Erste_Anpassung_Datum": format_date(anpassung),
            "Anzahlung_Prozent": format_percent(anzahlung) if anzahlung is not None else "",
            "Buchungen_Anzahl": str(int(buchungen_anzahl or 0)),
            "Erstes_Geschaeftsjahr": geschaeftsjahr,
        }
    )
    drops: list[str] = []
    for key, amount in pauschalen.items():
        if amount is None:
            drops.append(key)
            values[key] = ""
        else:
            values[key] = format_amount(amount)
    if not email:
        drops.append("Mandant_Email")

    blocks = {"Anzahlung": anzahlung is not None, "Verbraucher": bool(verbraucher)}
    summary = {
        "mandant_kurzname": kurzname,
        "vertragstyp": "Verbraucher" if verbraucher else "Unternehmer",
        "vertragsdatum": values["Vertragsdatum"],
        "erste_anpassung": values["Erste_Anpassung_Datum"],
        "anzahlung": f"{values['Anzahlung_Prozent']} %" if anzahlung is not None else "keine",
        "pauschalen": {
            key: (format_amount(amount) + " EUR") if amount is not None else "gestrichen" for key, amount in pauschalen.items()
        },
        "registrierung": "250,00 EUR einmalig (fest in der Vorlage)",
    }
    file_name = safe_file_name(f"Mandatsvertrag {kurzname} {jahr_text}")
    return {"values": values, "blocks": blocks, "drops": drops, "summary": summary, "file_name": file_name}


# ---------------------------------------------------------------------------
# Open-WebUI-Tool
# ---------------------------------------------------------------------------


class Tools:
    class Valves(BaseModel):
        TEMPLATE_PATH: str = Field(
            default=DEFAULT_TEMPLATE_PATH,
            description="Pfad zur Word-Vorlage im Container (Datenvolume). Die Vorlage liegt nicht im Repository.",
        )
        MAX_FIELD_CHARS: int = Field(default=500, description="Obergrenze je Textfeld in Zeichen.")

    def __init__(self) -> None:
        self.valves = self.Valves()

    async def beratungsvertrag_docx(
        self,
        mandant_name: str,
        mandant_anschrift: str,
        gespraechsdatum: str,
        mandant_email: str = "",
        mandant_kurzname: str = "",
        jahr: str = "",
        vertragsdatum: str = "",
        verbraucher: bool = False,
        anzahlung_prozent: str = "",
        buchungen_anzahl: int = 0,
        pauschale_buchfuehrung_monat: str = "",
        lohn_einrichtung: str = "",
        pauschale_jahresabschluss: str = "",
        erstes_geschaeftsjahr: str = "",
        pauschale_steuererklaerung: str = "",
        __user__: Optional[dict] = None,
        __event_emitter__: Any = None,
    ) -> str:
        """
        Erzeugt den Mandatsvertrag (Beratungsvertrag, Angebot über steuerliche Beratungsleistungen) als Word-Datei aus der Kanzleivorlage. Nutze dieses Tool immer, wenn ein Beratungsvertrag, Mandatsvertrag oder Angebot als Datei erstellt werden soll und die Stammdaten sowie die Pauschalen feststehen. Das Tool berechnet das Datum der ersten Honoraranpassung selbst, streicht nicht genutzte Pauschalzeilen und entfernt Bausteine für Anzahlung und Verbraucher, wenn sie nicht zutreffen. Schreibe den Vertragstext nicht selbst.

        :param mandant_name: Name bzw. Firma des Mandanten, z. B. "Beispiel Handel GmbH" oder "Erika Musterfrau".
        :param mandant_anschrift: Straße Hausnummer, PLZ Ort, in einer Zeile.
        :param gespraechsdatum: Datum des Erstgesprächs, Format TT.MM.JJJJ.
        :param mandant_email: E-Mail-Adresse des Mandanten. Leer = Zeile entfällt.
        :param mandant_kurzname: Kurzname für den Titel "Mandatsvertrag <Kurzname> <Jahr>". Leer = mandant_name.
        :param jahr: Jahr im Titel, JJJJ. Leer = Jahr des Vertragsdatums.
        :param vertragsdatum: Datum des Angebots, TT.MM.JJJJ. Leer = heute.
        :param verbraucher: true bei Privatpersonen (Verbraucher, Widerrufsbelehrung bleibt), false bei Unternehmern.
        :param anzahlung_prozent: Anzahlung in Prozent des Honorars, z. B. "30". Leer = keine Anzahlung, Baustein entfällt.
        :param buchungen_anzahl: Angenommene Buchungen je Monat für die Finanzbuchführungspauschale. Pflicht, wenn pauschale_buchfuehrung_monat gesetzt ist.
        :param pauschale_buchfuehrung_monat: Pauschale Finanzbuchführung je Monat netto in EUR, z. B. "450,00". Leer = Zeile wird gestrichen.
        :param lohn_einrichtung: Einmalige Einrichtungspauschale Lohnbuchführung netto in EUR. Leer = Zeile wird gestrichen.
        :param pauschale_jahresabschluss: Pauschale Jahresabschluss inkl. KSt/GewSt mit E-Bilanz netto in EUR. Leer = Zeile wird gestrichen.
        :param erstes_geschaeftsjahr: Erstes Geschäftsjahr der Jahresabschlusspauschale, JJJJ. Pflicht, wenn pauschale_jahresabschluss gesetzt ist.
        :param pauschale_steuererklaerung: Pauschale Steuererklärungen je Veranlagungszeitraum netto in EUR. Leer = Zeile wird gestrichen.
        :return: Ergebnisbericht mit Download-Link und den eingesetzten Werten. Gib dem Nutzer den Link unverändert weiter.
        """
        started = time.monotonic()
        if not __user__ or not __user__.get("id"):
            return self._fail("Kein Nutzerkontext vorhanden. Abgebrochen.")
        if not os.path.isfile(self.valves.TEMPLATE_PATH):
            return self._fail(
                "Die Word-Vorlage ist im Container nicht vorhanden. Der Admin legt sie unter dem Pfad aus der "
                "Einstellung TEMPLATE_PATH ab."
            )
        text_fields = {
            "mandant_name": mandant_name,
            "mandant_anschrift": mandant_anschrift,
            "mandant_email": mandant_email,
            "mandant_kurzname": mandant_kurzname,
        }
        too_long = [key for key, value in text_fields.items() if value and len(value) > self.valves.MAX_FIELD_CHARS]
        if too_long:
            return self._fail(f"Feld zu lang: {', '.join(too_long)} (erlaubt {self.valves.MAX_FIELD_CHARS} Zeichen).")

        try:
            contract = build_contract(
                mandant_name=mandant_name,
                mandant_anschrift=mandant_anschrift,
                gespraechsdatum=gespraechsdatum,
                mandant_email=mandant_email,
                mandant_kurzname=mandant_kurzname,
                jahr=jahr,
                vertragsdatum=vertragsdatum,
                verbraucher=verbraucher,
                anzahlung_prozent=anzahlung_prozent,
                buchungen_anzahl=buchungen_anzahl,
                pauschale_buchfuehrung_monat=pauschale_buchfuehrung_monat,
                lohn_einrichtung=lohn_einrichtung,
                pauschale_jahresabschluss=pauschale_jahresabschluss,
                erstes_geschaeftsjahr=erstes_geschaeftsjahr,
                pauschale_steuererklaerung=pauschale_steuererklaerung,
            )
        except ValueError as exc:
            return self._fail(f"Eingaben unvollständig oder ungültig: {exc}")

        await self._status(__event_emitter__, f"Befülle Vorlage, erzeuge {contract['file_name']} …")
        try:
            record = await self._fill_and_store(contract, __user__["id"])
        except ValueError as exc:
            return self._fail(str(exc))
        except Exception as exc:
            log.warning("beratungsvertrag_docx: fill/store failed (%s)", type(exc).__name__)
            return self._fail(f"Datei konnte nicht erzeugt werden: {type(exc).__name__}")

        url = f"/api/v1/files/{record['id']}/content"
        if __event_emitter__:
            await __event_emitter__(
                {
                    "type": "files",
                    "data": {
                        "files": [
                            {
                                "type": "file",
                                "id": record["id"],
                                "url": url,
                                "name": contract["file_name"],
                                "size": record["size"],
                                "content_type": DOCX_CONTENT_TYPE,
                            }
                        ]
                    },
                }
            )
        await self._status(__event_emitter__, "Vertrag fertig.", done=True)
        return json.dumps(
            {
                "status": "success",
                "message": f"Mandatsvertrag erstellt. Download: [{contract['file_name']}]({url})",
                "download_url": url,
                "file_name": contract["file_name"],
                "summary": contract["summary"],
                "blocks": record["stats"]["blocks"],
                "placeholders_replaced": record["stats"]["replaced"],
                "paragraphs_removed": record["stats"]["paragraphs_removed"],
                "seconds": round(time.monotonic() - started, 1),
                "instruction_for_assistant": (
                    "Zeige dem Nutzer den Download-Link als Markdown-Link und nenne kurz die eingesetzten Werte "
                    "aus summary. Weise auf offene Prüfpunkte hin, wiederhole den Vertragstext nicht."
                ),
            },
            ensure_ascii=False,
        )

    # ----------------------------------------------------------------- helpers

    @staticmethod
    def _fail(message: str) -> str:
        return json.dumps({"status": "error", "message": message}, ensure_ascii=False)

    @staticmethod
    async def _status(emitter: Any, description: str, done: bool = False) -> None:
        if not emitter:
            return
        try:
            await emitter({"type": "status", "data": {"description": description, "done": done}})
        except Exception:
            pass

    async def _fill_and_store(self, contract: dict, user_id: str) -> dict:
        from open_webui.models.files import FileForm, Files
        from open_webui.storage.provider import Storage

        with tempfile.TemporaryDirectory() as tmp:
            tmp_path = os.path.join(tmp, contract["file_name"])
            stats = await asyncio.to_thread(
                fill_docx,
                self.valves.TEMPLATE_PATH,
                contract["values"],
                contract["blocks"],
                contract["drops"],
                tmp_path,
            )
            if stats["remaining"]:
                raise ValueError(
                    "In der Vorlage blieben Platzhalter unbefüllt: " + ", ".join(stats["remaining"]) + ". "
                    "Vorlage und Tool passen nicht zusammen, bitte den Admin informieren."
                )
            with open(tmp_path, "rb") as fh:
                payload = fh.read()

        file_id = str(uuid.uuid4())
        tags = {"OpenWebUI-User-Id": user_id, "OpenWebUI-File-Id": file_id}
        contents, stored_path = await asyncio.to_thread(
            Storage.upload_file, io.BytesIO(payload), f"{file_id}_{contract['file_name']}", tags
        )
        record = await Files.insert_new_file(
            user_id,
            FileForm(
                id=file_id,
                filename=contract["file_name"],
                path=stored_path,
                data={},
                meta={
                    "name": contract["file_name"],
                    "content_type": DOCX_CONTENT_TYPE,
                    "size": len(contents),
                    "data": {"generated_by": "beratungsvertrag_docx"},
                },
            ),
        )
        if record is None:
            raise RuntimeError("Dateieintrag konnte nicht angelegt werden.")
        return {"id": file_id, "size": len(contents), "stats": stats}
