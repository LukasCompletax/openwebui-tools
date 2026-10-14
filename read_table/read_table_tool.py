"""
title: Tabelle lesen und auswerten (XLSX/CSV)
author: completax
version: 0.1.0
description: Liest eine im Chat angehängte Excel- oder CSV-Datei, beschreibt Blätter, Spalten und Kennzahlen und rechnet deterministisch Summen, Gruppierungen und Filter. Das Modell rechnet nicht selbst, es fragt das Tool.
requirements: openpyxl
"""

from __future__ import annotations

import asyncio
import csv
import io
import json
import logging
import os
import re
from dataclasses import dataclass, field
from decimal import Decimal, InvalidOperation
from typing import Any, Optional

from pydantic import BaseModel, Field

log = logging.getLogger(__name__)

SUPPORTED_INPUT = {".xlsx", ".xlsm", ".csv", ".tsv"}
OPERATIONS = {"sum", "count", "avg", "min", "max"}

# ---------------------------------------------------------------------------
# Kern ohne Open-WebUI-Abhängigkeit
# ---------------------------------------------------------------------------


@dataclass
class Table:
    name: str
    header: list[str]
    rows: list[list[Any]] = field(default_factory=list)

    @property
    def width(self) -> int:
        return len(self.header)


_NUMBER_DE = re.compile(r"^-?\d{1,3}(\.\d{3})*(,\d+)?$|^-?\d+(,\d+)?$")
_NUMBER_EN = re.compile(r"^-?\d{1,3}(,\d{3})*(\.\d+)?$|^-?\d+(\.\d+)?$")
_CURRENCY = re.compile(r"[€$£\s]|EUR|USD|CHF", re.IGNORECASE)


def parse_number(value: Any) -> Optional[Decimal]:
    """Zahl aus Zelle holen. Deutsche (1.234,56) und englische (1,234.56) Schreibweise.

    Gibt None zurück, wenn die Zelle keine Zahl ist. Geldbeträge werden als Decimal
    geführt, nie als float.
    """
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, Decimal):
        return value
    if isinstance(value, int):
        return Decimal(value)
    if isinstance(value, float):
        return Decimal(str(value))
    text = _CURRENCY.sub("", str(value)).strip()
    if not text:
        return None
    try:
        if _NUMBER_DE.match(text) and ("," in text or "." in text):
            return Decimal(text.replace(".", "").replace(",", "."))
        if _NUMBER_EN.match(text):
            return Decimal(text.replace(",", ""))
    except InvalidOperation:
        return None
    return None


def cell_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if hasattr(value, "isoformat"):
        return value.isoformat(sep=" ") if hasattr(value, "hour") else value.isoformat()
    return str(value)


def _normalize_header(raw: list[Any], width: int) -> list[str]:
    header: list[str] = []
    seen: dict[str, int] = {}
    for index in range(width):
        name = cell_text(raw[index]).strip() if index < len(raw) else ""
        name = name or f"Spalte {index + 1}"
        if name in seen:
            seen[name] += 1
            name = f"{name} ({seen[name]})"
        else:
            seen[name] = 1
        header.append(name)
    return header


def _finish_table(name: str, raw_rows: list[list[Any]], max_rows: int) -> Table:
    raw_rows = [row for row in raw_rows if any(cell not in (None, "") for cell in row)]
    if not raw_rows:
        return Table(name=name, header=[])
    width = max(len(row) for row in raw_rows)
    header = _normalize_header(raw_rows[0], width)
    body = [list(row) + [None] * (width - len(row)) for row in raw_rows[1 : max_rows + 1]]
    table = Table(name=name, header=header, rows=body)
    table.truncated = len(raw_rows) - 1 > max_rows  # type: ignore[attr-defined]
    table.total_rows = len(raw_rows) - 1  # type: ignore[attr-defined]
    return table


def load_xlsx(path: str, max_rows: int, sheet: str = "") -> list[Table]:
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    tables: list[Table] = []
    names = [sheet] if sheet else workbook.sheetnames
    for name in names:
        if name not in workbook.sheetnames:
            raise LookupError(f"Blatt '{name}' nicht gefunden. Vorhanden: {', '.join(workbook.sheetnames)}")
        rows = [list(row) for row in workbook[name].iter_rows(values_only=True, max_row=max_rows + 1)]
        # Für die Gesamtzeilenzahl ohne Begrenzung lesen, aber nur zählen
        total = sum(1 for row in workbook[name].iter_rows(values_only=True) if any(c not in (None, "") for c in row))
        table = _finish_table(name, rows, max_rows)
        table.total_rows = max(total - 1, 0)  # type: ignore[attr-defined]
        table.truncated = table.total_rows > len(table.rows)  # type: ignore[attr-defined]
        tables.append(table)
    workbook.close()
    return tables


def load_delimited(path: str, max_rows: int) -> list[Table]:
    raw: bytes = open(path, "rb").read()
    text = None
    for encoding in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            text = raw.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    if text is None:
        raise ValueError("Zeichenkodierung der Datei nicht erkannt.")
    sample = text[:4096]
    delimiter = ";"
    if path.lower().endswith(".tsv"):
        delimiter = "\t"
    else:
        try:
            delimiter = csv.Sniffer().sniff(sample, delimiters=";,\t|").delimiter
        except csv.Error:
            counts = {d: sample.count(d) for d in (";", ",", "\t", "|")}
            delimiter = max(counts, key=counts.get)  # type: ignore[arg-type]
    rows = [row for row in csv.reader(io.StringIO(text), delimiter=delimiter)]
    name = os.path.splitext(os.path.basename(path))[0]
    return [_finish_table(name, rows, max_rows)]


def load_tables(path: str, max_rows: int, sheet: str = "") -> list[Table]:
    ext = os.path.splitext(path)[1].lower()
    if ext in (".xlsx", ".xlsm"):
        return load_xlsx(path, max_rows, sheet)
    if ext in (".csv", ".tsv"):
        return load_delimited(path, max_rows)
    raise ValueError(f"Dateityp {ext or '(ohne Endung)'} wird nicht unterstützt.")


def column_index(table: Table, column: str) -> int:
    """Spalte per Name (Groß/Klein egal), per Buchstabe (A, B, ...) oder per Nummer (1-basiert)."""
    wanted = (column or "").strip()
    if not wanted:
        raise LookupError("Keine Spalte angegeben.")
    lowered = [h.lower() for h in table.header]
    if wanted.lower() in lowered:
        return lowered.index(wanted.lower())
    if wanted.isdigit() and 1 <= int(wanted) <= table.width:
        return int(wanted) - 1
    if re.fullmatch(r"[A-Za-z]{1,2}", wanted):
        index = 0
        for ch in wanted.upper():
            index = index * 26 + (ord(ch) - 64)
        if 1 <= index <= table.width:
            return index - 1
    partial = [i for i, h in enumerate(lowered) if wanted.lower() in h]
    if len(partial) == 1:
        return partial[0]
    raise LookupError(f"Spalte '{column}' nicht gefunden. Vorhanden: {', '.join(table.header)}")


def describe_column(table: Table, index: int) -> dict:
    values = [row[index] for row in table.rows]
    numbers = [n for n in (parse_number(v) for v in values) if n is not None]
    empty = sum(1 for v in values if v in (None, ""))
    info: dict[str, Any] = {
        "name": table.header[index],
        "filled": len(values) - empty,
        "empty": empty,
    }
    if numbers and len(numbers) >= max(1, (len(values) - empty) * 0.8):
        info["type"] = "number"
        info["sum"] = str(sum(numbers))
        info["min"] = str(min(numbers))
        info["max"] = str(max(numbers))
    else:
        distinct = {cell_text(v) for v in values if v not in (None, "")}
        info["type"] = "text"
        info["distinct"] = len(distinct)
        if 0 < len(distinct) <= 12:
            info["values"] = sorted(distinct)[:12]
    return info


def describe_table(table: Table, preview_rows: int) -> dict:
    return {
        "sheet": table.name,
        "rows": getattr(table, "total_rows", len(table.rows)),
        "rows_loaded": len(table.rows),
        "truncated": getattr(table, "truncated", False),
        "columns": [describe_column(table, i) for i in range(table.width)],
        "preview": markdown_preview(table, preview_rows),
    }


def markdown_preview(table: Table, limit: int) -> str:
    if not table.header:
        return "(leer)"
    lines = ["| " + " | ".join(h.replace("|", "\\|") for h in table.header) + " |"]
    lines.append("|" + "---|" * table.width)
    for row in table.rows[:limit]:
        lines.append("| " + " | ".join(cell_text(c).replace("|", "\\|") for c in row) + " |")
    if len(table.rows) > limit:
        lines.append(f"… {len(table.rows) - limit} weitere Zeilen")
    return "\n".join(lines)


def _matches(cell: Any, wanted: str) -> bool:
    text = cell_text(cell).strip().lower()
    wanted = wanted.strip().lower()
    number, wanted_number = parse_number(cell), parse_number(wanted)
    if number is not None and wanted_number is not None:
        return number == wanted_number
    return text == wanted or (bool(wanted) and wanted in text)


def aggregate(
    table: Table,
    value_column: str,
    operation: str,
    group_by: str = "",
    filter_column: str = "",
    filter_value: str = "",
    max_groups: int = 200,
) -> dict:
    """Summe, Anzahl, Durchschnitt, Minimum oder Maximum, optional gruppiert und gefiltert.

    Rechnet mit Decimal. Zellen ohne Zahl werden bei sum/avg/min/max ignoriert und gezählt.
    """
    operation = (operation or "sum").strip().lower()
    if operation not in OPERATIONS:
        raise ValueError(f"Operation '{operation}' unbekannt. Erlaubt: {', '.join(sorted(OPERATIONS))}")
    value_index = column_index(table, value_column) if (value_column or operation != "count") else None
    group_index = column_index(table, group_by) if group_by else None
    filter_index = column_index(table, filter_column) if filter_column else None

    rows = table.rows
    if filter_index is not None:
        rows = [r for r in rows if _matches(r[filter_index], filter_value)]

    groups: dict[str, list[Any]] = {}
    for row in rows:
        key = cell_text(row[group_index]).strip() if group_index is not None else "__all__"
        groups.setdefault(key or "(leer)", []).append(row)

    def compute(group_rows: list[list[Any]]) -> dict:
        if operation == "count" and value_index is None:
            return {"count": len(group_rows)}
        numbers = [n for n in (parse_number(r[value_index]) for r in group_rows) if n is not None]  # type: ignore[index]
        skipped = len(group_rows) - len(numbers)
        result: dict[str, Any] = {"count": len(numbers), "non_numeric_skipped": skipped}
        if operation == "count":
            return result
        if not numbers:
            result["value"] = None
            return result
        if operation == "sum":
            value = sum(numbers)
        elif operation == "avg":
            value = (sum(numbers) / len(numbers)).quantize(Decimal("0.0001"))
        elif operation == "min":
            value = min(numbers)
        else:
            value = max(numbers)
        result["value"] = str(value)
        return result

    if group_index is None:
        return {
            "operation": operation,
            "column": table.header[value_index] if value_index is not None else None,
            "rows_considered": len(rows),
            **compute(rows),
        }
    items = sorted(groups.items(), key=lambda kv: kv[0])
    result_groups = [{"group": key, **compute(group_rows)} for key, group_rows in items[:max_groups]]
    return {
        "operation": operation,
        "column": table.header[value_index] if value_index is not None else None,
        "group_by": table.header[group_index],
        "rows_considered": len(rows),
        "groups": result_groups,
        "groups_truncated": len(items) > max_groups,
    }


# ---------------------------------------------------------------------------
# Open-WebUI-Tool
# ---------------------------------------------------------------------------


class Tools:
    class Valves(BaseModel):
        MAX_ROWS: int = Field(default=5000, description="Höchstens so viele Datenzeilen je Blatt laden.")
        PREVIEW_ROWS: int = Field(default=15, description="Zeilen in der Markdown-Vorschau für das Modell.")
        MAX_SHEETS: int = Field(default=10, description="Höchstens so viele Blätter beschreiben.")
        MAX_GROUPS: int = Field(default=200, description="Höchstens so viele Gruppen bei Gruppierungen zurückgeben.")

    def __init__(self) -> None:
        self.valves = self.Valves()

    async def describe_spreadsheet(
        self,
        file_name: str = "",
        sheet: str = "",
        __user__: Optional[dict] = None,
        __files__: Optional[list] = None,
        __event_emitter__: Any = None,
    ) -> str:
        """
        Beschreibt eine im Chat angehängte Excel- oder CSV-Datei: Blätter, Zeilenzahl, Spalten mit Typ, Summe, Minimum, Maximum oder Werteliste, plus eine Vorschau der ersten Zeilen als Markdown-Tabelle. Nutze dieses Tool zuerst, wenn der Nutzer Fragen zu einer angehängten Tabelle stellt. Rechne danach nicht selbst, sondern nutze aggregate_spreadsheet.

        :param file_name: Name der angehängten Datei, falls mehrere angehängt sind. Leer = zuletzt angehängte Tabelle.
        :param sheet: Nur dieses Blatt beschreiben. Leer = alle Blätter.
        :return: JSON mit Blättern, Spaltenbeschreibung und Vorschau.
        """
        try:
            source = await self._resolve_source_file(__files__ or [], file_name, (__user__ or {}).get("id"))
            await self._status(__event_emitter__, f"Lese {source['name']} …")
            tables = await asyncio.to_thread(load_tables, source["path"], self.valves.MAX_ROWS, sheet)
        except (LookupError, ValueError) as exc:
            return self._fail(str(exc))
        except Exception as exc:
            log.warning("read_table: describe failed (%s)", type(exc).__name__)
            return self._fail(f"Datei konnte nicht gelesen werden: {type(exc).__name__}")

        sheets = [describe_table(t, self.valves.PREVIEW_ROWS) for t in tables[: self.valves.MAX_SHEETS]]
        await self._status(__event_emitter__, "Tabelle gelesen.", done=True)
        return json.dumps(
            {
                "status": "success",
                "file": source["name"],
                "sheets": sheets,
                "sheets_truncated": len(tables) > self.valves.MAX_SHEETS,
                "instruction_for_assistant": (
                    "Für Summen, Gruppierungen oder Filter aggregate_spreadsheet aufrufen, nicht selbst rechnen. "
                    "Beträge sind exakte Dezimalzahlen als Text."
                ),
            },
            ensure_ascii=False,
        )

    async def aggregate_spreadsheet(
        self,
        value_column: str = "",
        operation: str = "sum",
        group_by: str = "",
        filter_column: str = "",
        filter_value: str = "",
        sheet: str = "",
        file_name: str = "",
        __user__: Optional[dict] = None,
        __files__: Optional[list] = None,
        __event_emitter__: Any = None,
    ) -> str:
        """
        Rechnet deterministisch auf einer angehängten Excel- oder CSV-Datei: Summe, Anzahl, Durchschnitt, Minimum oder Maximum einer Spalte, optional gruppiert nach einer anderen Spalte und optional gefiltert. Nutze dieses Tool für jede Zahl, die der Nutzer aus der Tabelle wissen will.

        :param value_column: Spalte mit den Werten, per Name, Buchstabe (B) oder Nummer. Bei operation=count optional.
        :param operation: sum, count, avg, min oder max. Standard sum.
        :param group_by: Spalte, nach der gruppiert wird, z. B. "Konto" oder "Monat". Leer = keine Gruppierung.
        :param filter_column: Spalte für einen Filter. Leer = kein Filter.
        :param filter_value: Wert, den filter_column haben muss (exakt oder enthalten, Zahlen werden verglichen).
        :param sheet: Blattname bei Excel. Leer = erstes Blatt.
        :param file_name: Name der angehängten Datei, falls mehrere angehängt sind.
        :return: JSON mit Ergebnis, bei Gruppierung je Gruppe. Werte sind exakte Dezimalzahlen als Text.
        """
        try:
            source = await self._resolve_source_file(__files__ or [], file_name, (__user__ or {}).get("id"))
            await self._status(__event_emitter__, f"Rechne auf {source['name']} …")
            tables = await asyncio.to_thread(load_tables, source["path"], self.valves.MAX_ROWS, sheet)
            table = tables[0]
            if not table.header:
                return self._fail("Das Blatt ist leer.")
            result = aggregate(
                table, value_column, operation, group_by, filter_column, filter_value, self.valves.MAX_GROUPS
            )
        except (LookupError, ValueError) as exc:
            return self._fail(str(exc))
        except Exception as exc:
            log.warning("read_table: aggregate failed (%s)", type(exc).__name__)
            return self._fail(f"Auswertung fehlgeschlagen: {type(exc).__name__}")

        await self._status(__event_emitter__, "Auswertung fertig.", done=True)
        return json.dumps(
            {
                "status": "success",
                "file": source["name"],
                "sheet": table.name,
                "rows_loaded": len(table.rows),
                "truncated": getattr(table, "truncated", False),
                **result,
                "instruction_for_assistant": "Zahlen unverändert übernehmen, nicht runden oder neu berechnen.",
            },
            ensure_ascii=False,
        )

    # ----------------------------------------------------------------- helpers

    @staticmethod
    def _fail(message: str) -> str:
        return json.dumps({"status": "error", "message": message}, ensure_ascii=False)

    @staticmethod
    async def _status(emitter: Any, description: str, done: bool = False) -> None:
        if not emitter:
            return
        try:
            await emitter({"type": "status", "data": {"description": description, "done": done}})
        except Exception:
            pass

    async def _resolve_source_file(self, files: list, wanted_name: str, user_id: Optional[str]) -> dict:
        from open_webui.models.files import Files
        from open_webui.storage.provider import Storage

        if not user_id:
            raise LookupError("Kein Nutzerkontext vorhanden. Abgebrochen.")
        candidates = []
        for item in files:
            if not isinstance(item, dict) or item.get("type") != "file":
                continue
            name = item.get("name") or (item.get("file") or {}).get("filename") or ""
            if os.path.splitext(name)[1].lower() in SUPPORTED_INPUT and item.get("id"):
                candidates.append({"id": item["id"], "name": name})
        if not candidates:
            raise LookupError("Keine Excel- oder CSV-Datei im Chat angehängt (xlsx, xlsm, csv, tsv).")
        if wanted_name:
            wanted = wanted_name.strip().lower()
            matches = [c for c in candidates if c["name"].lower() == wanted] or [
                c for c in candidates if wanted in c["name"].lower()
            ]
            if not matches:
                raise LookupError(f"Datei '{wanted_name}' nicht gefunden. Angehängt: {', '.join(c['name'] for c in candidates)}")
            chosen = matches[-1]
        else:
            chosen = candidates[-1]
        record = await Files.get_file_by_id(chosen["id"])
        if record is None or record.user_id != user_id:
            raise LookupError("Zugriff auf die angehängte Datei verweigert.")
        if not record.path:
            raise LookupError("Die Datei hat keinen Speicherpfad. Bitte erneut hochladen.")
        local_path = await asyncio.to_thread(Storage.get_file, record.path)
        if not local_path or not os.path.exists(local_path):
            raise LookupError("Die Datei ist auf dem Server nicht mehr vorhanden.")
        return {"id": record.id, "name": chosen["name"] or record.filename, "path": local_path}
