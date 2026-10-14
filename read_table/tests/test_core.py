"""Tests für den Kern von read_table, ohne Open WebUI. Alle Inhalte sind erfunden.

Ausführen:
    python read_table/tests/test_core.py
"""

from __future__ import annotations

import os
import sys
import tempfile
from decimal import Decimal

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))

import read_table_tool as t  # noqa: E402

ROWS = [
    ["Datum", "Konto", "Buchungstext", "Betrag", "Belegnummer"],
    ["2026-01-05", "4400", "Erlöse Beratung", 1200.5, "RE-1"],
    ["2026-01-09", "4400", "Erlöse Beratung", 800, "RE-2"],
    ["2026-01-12", "6300", "Büromaterial", "-45,90", "ER-7"],
    ["2026-02-02", "6300", "Fachliteratur", "-120,00 €", "ER-8"],
    ["2026-02-03", "4400", "Erlöse Schulung", "1.000,00", "RE-3"],
    ["2026-02-10", "6800", "Porto", None, "ER-9"],
]


def make_xlsx(path: str) -> None:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Buchungen"
    for row in ROWS:
        ws.append(row)
    ws2 = wb.create_sheet("Leer")
    ws3 = wb.create_sheet("Konten")
    ws3.append(["Konto", "Bezeichnung"])
    ws3.append([4400, "Erlöse"])
    wb.save(path)


def make_csv(path: str, delimiter: str = ";", encoding: str = "utf-8-sig") -> None:
    import csv

    with open(path, "w", encoding=encoding, newline="") as fh:
        writer = csv.writer(fh, delimiter=delimiter)
        for row in ROWS:
            writer.writerow(["" if c is None else c for c in row])


def test_parse_number_variants():
    assert t.parse_number("1.234,56") == Decimal("1234.56")
    assert t.parse_number("1,234.56") == Decimal("1234.56")
    assert t.parse_number("-45,90") == Decimal("-45.90")
    assert t.parse_number("-120,00 €") == Decimal("-120.00")
    assert t.parse_number("1200,5 EUR") == Decimal("1200.5")
    assert t.parse_number(800) == Decimal("800")
    assert t.parse_number(0.1) == Decimal("0.1"), "float über str, kein Binärrauschen"
    assert t.parse_number("RE-1") is None
    assert t.parse_number("") is None and t.parse_number(None) is None
    assert t.parse_number(True) is None


def test_load_xlsx_sheets_and_truncation():
    with tempfile.TemporaryDirectory() as tmp:
        path = os.path.join(tmp, "b.xlsx")
        make_xlsx(path)
        tables = t.load_tables(path, max_rows=100)
        assert [tb.name for tb in tables] == ["Buchungen", "Leer", "Konten"]
        buch = tables[0]
        assert buch.header == ["Datum", "Konto", "Buchungstext", "Betrag", "Belegnummer"]
        assert len(buch.rows) == 6 and buch.total_rows == 6 and buch.truncated is False
        assert tables[1].header == [] and tables[1].rows == []

        small = t.load_tables(path, max_rows=2, sheet="Buchungen")[0]
        assert len(small.rows) == 2 and small.total_rows == 6 and small.truncated is True

        try:
            t.load_tables(path, 10, sheet="Gibtsnicht")
        except LookupError as exc:
            assert "Buchungen" in str(exc)
        else:
            raise AssertionError("expected LookupError")


def test_load_csv_semicolon_and_comma_and_cp1252():
    with tempfile.TemporaryDirectory() as tmp:
        p1 = os.path.join(tmp, "a.csv")
        make_csv(p1, ";")
        table = t.load_tables(p1, 100)[0]
        assert table.header[3] == "Betrag" and len(table.rows) == 6
        p2 = os.path.join(tmp, "b.csv")
        make_csv(p2, ",")
        assert t.load_tables(p2, 100)[0].header == ROWS[0]
        p3 = os.path.join(tmp, "c.csv")
        make_csv(p3, ";", encoding="cp1252")
        assert "Erlöse Beratung" in [r[2] for r in t.load_tables(p3, 100)[0].rows]


def test_column_index_by_name_letter_number_and_partial():
    table = t.Table("x", ["Datum", "Konto", "Betrag"], [])
    assert t.column_index(table, "betrag") == 2
    assert t.column_index(table, "B") == 1
    assert t.column_index(table, "3") == 2
    assert t.column_index(table, "Dat") == 0
    try:
        t.column_index(table, "Umsatz")
    except LookupError as exc:
        assert "Vorhanden" in str(exc)
    else:
        raise AssertionError("expected LookupError")


def _buchungen() -> t.Table:
    with tempfile.TemporaryDirectory() as tmp:
        path = os.path.join(tmp, "b.xlsx")
        make_xlsx(path)
        return t.load_tables(path, 100, sheet="Buchungen")[0]


def test_describe_columns():
    table = _buchungen()
    desc = t.describe_table(table, preview_rows=3)
    cols = {c["name"]: c for c in desc["columns"]}
    assert cols["Betrag"]["type"] == "number"
    assert cols["Betrag"]["sum"] == "2834.60"
    assert cols["Betrag"]["min"] == "-120.00" and cols["Betrag"]["max"] == "1200.5"
    assert cols["Betrag"]["empty"] == 1
    assert cols["Buchungstext"]["type"] == "text" and cols["Buchungstext"]["distinct"] == 5
    assert "Erlöse Beratung" in cols["Buchungstext"]["values"]
    assert desc["preview"].startswith("| Datum | Konto |")
    assert "weitere Zeilen" in desc["preview"]


def test_aggregate_sum_group_filter():
    table = _buchungen()
    total = t.aggregate(table, "Betrag", "sum")
    assert total["value"] == "2834.60" and total["count"] == 5 and total["non_numeric_skipped"] == 1

    by_konto = t.aggregate(table, "Betrag", "sum", group_by="Konto")
    groups = {g["group"]: g for g in by_konto["groups"]}
    assert groups["4400"]["value"] == "3000.50"
    assert groups["6300"]["value"] == "-165.90"
    assert groups["6800"]["value"] is None and groups["6800"]["non_numeric_skipped"] == 1

    filtered = t.aggregate(table, "Betrag", "sum", filter_column="Konto", filter_value="4400")
    assert filtered["value"] == "3000.50" and filtered["rows_considered"] == 3

    contains = t.aggregate(table, "Betrag", "count", filter_column="Buchungstext", filter_value="erlöse")
    assert contains["count"] == 3

    rows_only = t.aggregate(table, "", "count")
    assert rows_only["count"] == 6

    avg = t.aggregate(table, "Betrag", "avg", filter_column="Konto", filter_value="4400")
    assert avg["value"] == "1000.1667"
    assert t.aggregate(table, "Betrag", "max")["value"] == "1200.5"

    try:
        t.aggregate(table, "Betrag", "median")
    except ValueError as exc:
        assert "Operation" in str(exc)
    else:
        raise AssertionError("expected ValueError")


def test_aggregate_uses_decimal_not_float():
    table = t.Table("x", ["Betrag"], [["0,1"], ["0,2"]])
    assert t.aggregate(table, "Betrag", "sum")["value"] == "0.3"


if __name__ == "__main__":
    failures = 0
    for name, fn in sorted(globals().items()):
        if name.startswith("test_") and callable(fn):
            try:
                fn()
                print(f"PASS {name}")
            except Exception as exc:  # noqa: BLE001
                failures += 1
                print(f"FAIL {name}: {type(exc).__name__}: {exc}")
    sys.exit(1 if failures else 0)
