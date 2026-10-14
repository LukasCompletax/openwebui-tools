"""
title: Datei erzeugen (DOCX/PDF/XLSX/MD/CSV aus Markdown)
author: completax
version: 0.1.0
description: Erzeugt aus Markdown-Inhalt eine Datei zum Download: Word (DOCX), PDF, Excel (XLSX), Markdown, Text oder CSV. Das Modell schreibt den Inhalt, das Tool rendert deterministisch.
requirements: reportlab, python-docx, openpyxl
"""

from __future__ import annotations

import asyncio
import csv
import io
import json
import logging
import os
import re
import tempfile
import time
import uuid
from dataclasses import dataclass, field
from typing import Any, Optional
from xml.sax.saxutils import escape as xml_escape

from pydantic import BaseModel, Field

log = logging.getLogger(__name__)

SUPPORTED_OUTPUT = {"docx", "pdf", "xlsx", "md", "txt", "csv"}
CONTENT_TYPES = {
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "pdf": "application/pdf",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "md": "text/markdown",
    "txt": "text/plain",
    "csv": "text/csv",
}
DEFAULT_FONT_PATH = "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"
DEFAULT_BOLD_FONT_PATH = "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"

# ---------------------------------------------------------------------------
# Markdown → Blockstruktur (bewusst klein gehalten, keine Abhängigkeit)
# ---------------------------------------------------------------------------


@dataclass
class Run:
    text: str
    bold: bool = False
    italic: bool = False
    code: bool = False


@dataclass
class Node:
    kind: str  # heading | paragraph | bullets | numbers | table | code | rule
    text: str = ""
    level: int = 0
    items: list[str] = field(default_factory=list)
    rows: list[list[str]] = field(default_factory=list)
    header: bool = False


_HEADING = re.compile(r"^(#{1,6})\s+(.*?)\s*#*\s*$")
_BULLET = re.compile(r"^\s{0,3}[-*+]\s+(.*)$")
_NUMBER = re.compile(r"^\s{0,3}\d+[.)]\s+(.*)$")
_TABLE_SEP = re.compile(r"^\s*\|?\s*:?-{2,}:?\s*(\|\s*:?-{2,}:?\s*)*\|?\s*$")
_RULE = re.compile(r"^\s*([-*_])(\s*\1){2,}\s*$")
_FENCE = re.compile(r"^\s*```")
_INLINE = re.compile(r"(\*\*(.+?)\*\*|__(.+?)__|`([^`]+)`|\*(?!\s)(.+?)(?<!\s)\*|_(?!\s)(.+?)(?<!\s)_)")


def split_table_row(line: str) -> list[str]:
    stripped = line.strip()
    if stripped.startswith("|"):
        stripped = stripped[1:]
    if stripped.endswith("|"):
        stripped = stripped[:-1]
    return [cell.strip().replace("\\|", "|") for cell in re.split(r"(?<!\\)\|", stripped)]


def parse_inline(text: str) -> list[Run]:
    """**fett**, *kursiv*, `code` in Runs zerlegen. Alles andere bleibt Text."""
    runs: list[Run] = []
    pos = 0
    for match in _INLINE.finditer(text):
        if match.start() > pos:
            runs.append(Run(text[pos : match.start()]))
        bold = match.group(2) or match.group(3)
        code = match.group(4)
        italic = match.group(5) or match.group(6)
        if bold is not None:
            runs.append(Run(bold, bold=True))
        elif code is not None:
            runs.append(Run(code, code=True))
        else:
            runs.append(Run(italic or "", italic=True))
        pos = match.end()
    if pos < len(text):
        runs.append(Run(text[pos:]))
    return runs or [Run("")]


def plain_text(text: str) -> str:
    return "".join(run.text for run in parse_inline(text))


def parse_markdown(text: str) -> list[Node]:
    """Markdown in Blöcke zerlegen: Überschriften, Absätze, Listen, Tabellen, Code, Linien."""
    lines = text.replace("\r\n", "\n").replace("\r", "\n").split("\n")
    nodes: list[Node] = []
    paragraph: list[str] = []
    i = 0

    def flush_paragraph() -> None:
        if paragraph:
            nodes.append(Node("paragraph", text=" ".join(s.strip() for s in paragraph)))
            paragraph.clear()

    while i < len(lines):
        line = lines[i]
        if _FENCE.match(line):
            flush_paragraph()
            code_lines: list[str] = []
            i += 1
            while i < len(lines) and not _FENCE.match(lines[i]):
                code_lines.append(lines[i])
                i += 1
            nodes.append(Node("code", text="\n".join(code_lines)))
            i += 1
            continue
        if not line.strip():
            flush_paragraph()
            i += 1
            continue
        heading = _HEADING.match(line)
        if heading:
            flush_paragraph()
            nodes.append(Node("heading", text=heading.group(2), level=len(heading.group(1))))
            i += 1
            continue
        if _RULE.match(line):
            flush_paragraph()
            nodes.append(Node("rule"))
            i += 1
            continue
        if "|" in line and i + 1 < len(lines) and _TABLE_SEP.match(lines[i + 1]):
            flush_paragraph()
            rows = [split_table_row(line)]
            i += 2
            while i < len(lines) and "|" in lines[i] and lines[i].strip():
                rows.append(split_table_row(lines[i]))
                i += 1
            width = max(len(r) for r in rows)
            rows = [r + [""] * (width - len(r)) for r in rows]
            nodes.append(Node("table", rows=rows, header=True))
            continue
        if _BULLET.match(line) or _NUMBER.match(line):
            flush_paragraph()
            kind = "bullets" if _BULLET.match(line) else "numbers"
            pattern = _BULLET if kind == "bullets" else _NUMBER
            items: list[str] = []
            while i < len(lines) and pattern.match(lines[i]):
                items.append(pattern.match(lines[i]).group(1).strip())  # type: ignore[union-attr]
                i += 1
                # Fortsetzungszeilen, eingerückt
                while i < len(lines) and lines[i].startswith("  ") and lines[i].strip() and not pattern.match(lines[i]):
                    items[-1] += " " + lines[i].strip()
                    i += 1
            nodes.append(Node(kind, items=items))
            continue
        paragraph.append(line)
        i += 1
    flush_paragraph()
    return nodes


def extract_tables(nodes: list[Node]) -> list[Node]:
    return [n for n in nodes if n.kind == "table"]


def parse_delimited(text: str) -> Optional[list[list[str]]]:
    """CSV-artigen Text erkennen (Komma, Semikolon oder Tab). None, wenn keine Tabelle."""
    lines = [line for line in text.strip().splitlines() if line.strip()]
    if len(lines) < 2:
        return None
    for delimiter in (";", "\t", ","):
        counts = [line.count(delimiter) for line in lines]
        if counts[0] >= 1 and all(c == counts[0] for c in counts):
            return [row for row in csv.reader(lines, delimiter=delimiter)]
    return None


_NUMERIC = re.compile(r"^-?\d+(\.\d+)?$")


def cell_value(cell: str) -> Any:
    """Nur eindeutig maschinenlesbare Zahlen umwandeln. Deutsche Formate bleiben Text."""
    stripped = cell.strip()
    if _NUMERIC.match(stripped):
        return int(stripped) if "." not in stripped else float(stripped)
    return plain_text(stripped)


# ---------------------------------------------------------------------------
# Renderer
# ---------------------------------------------------------------------------


def render_markdown(nodes: list[Node], original: str) -> str:
    return original.strip() + "\n"


def render_text(nodes: list[Node]) -> str:
    out: list[str] = []
    for node in nodes:
        if node.kind == "heading":
            title = plain_text(node.text)
            out.append(title)
            out.append(("=" if node.level == 1 else "-") * len(title))
        elif node.kind == "paragraph":
            out.append(plain_text(node.text))
        elif node.kind == "bullets":
            out.extend(f"- {plain_text(item)}" for item in node.items)
        elif node.kind == "numbers":
            out.extend(f"{n}. {plain_text(item)}" for n, item in enumerate(node.items, 1))
        elif node.kind == "table":
            widths = [max(len(plain_text(r[c])) for r in node.rows) for c in range(len(node.rows[0]))]
            for r_index, row in enumerate(node.rows):
                out.append("  ".join(plain_text(cell).ljust(widths[c]) for c, cell in enumerate(row)))
                if r_index == 0 and node.header:
                    out.append("  ".join("-" * w for w in widths))
        elif node.kind == "code":
            out.append(node.text)
        elif node.kind == "rule":
            out.append("-" * 40)
        out.append("")
    return "\n".join(out).rstrip() + "\n"


def render_csv(rows: list[list[str]], delimiter: str) -> str:
    buffer = io.StringIO()
    writer = csv.writer(buffer, delimiter=delimiter, lineterminator="\n")
    for row in rows:
        writer.writerow([plain_text(cell) for cell in row])
    return buffer.getvalue()


def render_xlsx(tables: list[tuple[str, list[list[str]]]], out_path: str) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    workbook = Workbook()
    workbook.remove(workbook.active)
    used_names: set[str] = set()
    for index, (name, rows) in enumerate(tables, start=1):
        sheet_name = re.sub(r"[\[\]:*?/\\]", " ", name or f"Tabelle {index}").strip()[:28] or f"Tabelle {index}"
        base, n = sheet_name, 2
        while sheet_name in used_names:
            sheet_name, n = f"{base} {n}", n + 1
        used_names.add(sheet_name)
        sheet = workbook.create_sheet(sheet_name)
        for row in rows:
            sheet.append([cell_value(cell) for cell in row])
        if rows:
            for cell in sheet[1]:
                cell.font = Font(bold=True)
            for col in range(1, len(rows[0]) + 1):
                width = max(len(str(r[col - 1])) if col - 1 < len(r) else 0 for r in rows)
                sheet.column_dimensions[get_column_letter(col)].width = min(max(10, width + 2), 60)
            sheet.freeze_panes = "A2"
    workbook.save(out_path)


def render_docx(nodes: list[Node], out_path: str, title: str) -> None:
    import docx
    from docx.shared import Pt

    document = docx.Document()
    if title:
        document.core_properties.title = title

    def add_runs(paragraph: Any, text: str) -> None:
        for run in parse_inline(text):
            r = paragraph.add_run(run.text)
            r.bold = run.bold or None
            r.italic = run.italic or None
            if run.code:
                r.font.name = "Courier New"

    for node in nodes:
        if node.kind == "heading":
            heading = document.add_heading(level=min(node.level, 4))
            add_runs(heading, node.text)
        elif node.kind == "paragraph":
            add_runs(document.add_paragraph(), node.text)
        elif node.kind == "bullets":
            for item in node.items:
                add_runs(document.add_paragraph(style="List Bullet"), item)
        elif node.kind == "numbers":
            for item in node.items:
                add_runs(document.add_paragraph(style="List Number"), item)
        elif node.kind == "table":
            table = document.add_table(rows=len(node.rows), cols=len(node.rows[0]))
            table.style = "Table Grid"
            for r_index, row in enumerate(node.rows):
                for c_index, cell in enumerate(row):
                    paragraph = table.cell(r_index, c_index).paragraphs[0]
                    add_runs(paragraph, cell)
                    if r_index == 0 and node.header:
                        for run in paragraph.runs:
                            run.bold = True
            document.add_paragraph()
        elif node.kind == "code":
            paragraph = document.add_paragraph()
            run = paragraph.add_run(node.text)
            run.font.name = "Courier New"
            run.font.size = Pt(9)
        elif node.kind == "rule":
            document.add_paragraph("―" * 30)
    document.save(out_path)


def _register_fonts(font_path: str, bold_path: str) -> tuple[str, str]:
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    if font_path and os.path.exists(font_path):
        if "DocFont" not in pdfmetrics.getRegisteredFontNames():
            pdfmetrics.registerFont(TTFont("DocFont", font_path))
        bold = "DocFont"
        if bold_path and os.path.exists(bold_path):
            if "DocFont-Bold" not in pdfmetrics.getRegisteredFontNames():
                pdfmetrics.registerFont(TTFont("DocFont-Bold", bold_path))
            bold = "DocFont-Bold"
        from reportlab.pdfbase.pdfmetrics import registerFontFamily

        registerFontFamily("DocFont", normal="DocFont", bold=bold, italic="DocFont", boldItalic=bold)
        return "DocFont", bold
    return "Helvetica", "Helvetica-Bold"


def _inline_markup(text: str) -> str:
    parts: list[str] = []
    for run in parse_inline(text):
        chunk = xml_escape(run.text)
        if run.code:
            chunk = f'<font face="Courier">{chunk}</font>'
        if run.bold:
            chunk = f"<b>{chunk}</b>"
        if run.italic:
            chunk = f"<i>{chunk}</i>"
        parts.append(chunk)
    return "".join(parts)


def render_pdf(
    nodes: list[Node],
    out_path: str,
    title: str,
    font_path: str = DEFAULT_FONT_PATH,
    bold_font_path: str = DEFAULT_BOLD_FONT_PATH,
    base_size: float = 10.5,
) -> None:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import (
        HRFlowable,
        ListFlowable,
        ListItem,
        Paragraph,
        Preformatted,
        SimpleDocTemplate,
        Spacer,
        Table,
        TableStyle,
    )

    font, bold = _register_fonts(font_path, bold_font_path)
    body = ParagraphStyle("Body", fontName=font, fontSize=base_size, leading=base_size * 1.35, spaceAfter=6)
    cell_style = ParagraphStyle("Cell", parent=body, fontSize=base_size - 1, leading=(base_size - 1) * 1.3, spaceAfter=0)
    code_style = ParagraphStyle("Code", fontName="Courier", fontSize=base_size - 1.5, leading=(base_size - 1.5) * 1.3, spaceAfter=6)
    headings = {
        level: ParagraphStyle(
            f"H{level}",
            fontName=bold,
            fontSize=base_size + max(0, 8 - 2 * level),
            leading=(base_size + max(0, 8 - 2 * level)) * 1.3,
            spaceBefore=10 if level <= 2 else 6,
            spaceAfter=6,
        )
        for level in range(1, 7)
    }
    doc = SimpleDocTemplate(
        out_path,
        pagesize=A4,
        leftMargin=22 * mm,
        rightMargin=22 * mm,
        topMargin=20 * mm,
        bottomMargin=20 * mm,
        title=title,
    )
    story: list[Any] = []
    for node in nodes:
        if node.kind == "heading":
            story.append(Paragraph(_inline_markup(node.text), headings[min(node.level, 6)]))
        elif node.kind == "paragraph":
            story.append(Paragraph(_inline_markup(node.text), body))
        elif node.kind in ("bullets", "numbers"):
            items = [ListItem(Paragraph(_inline_markup(item), body), leftIndent=12) for item in node.items]
            story.append(
                ListFlowable(
                    items,
                    bulletType="1" if node.kind == "numbers" else "bullet",
                    bulletFormat="%s." if node.kind == "numbers" else None,
                    bulletFontName=font,
                    bulletFontSize=base_size,
                    leftIndent=14,
                )
            )
        elif node.kind == "table":
            data = [[Paragraph(_inline_markup(cell), cell_style) for cell in row] for row in node.rows]
            table = Table(data, repeatRows=1 if node.header else 0, hAlign="LEFT")
            style = [
                ("GRID", (0, 0), (-1, -1), 0.4, colors.grey),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LEFTPADDING", (0, 0), (-1, -1), 4),
                ("RIGHTPADDING", (0, 0), (-1, -1), 4),
            ]
            if node.header:
                style.append(("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E8E8E8")))
            table.setStyle(TableStyle(style))
            story.append(table)
            story.append(Spacer(1, 6))
        elif node.kind == "code":
            story.append(Preformatted(node.text, code_style))
        elif node.kind == "rule":
            story.append(HRFlowable(width="100%", thickness=0.6, color=colors.grey, spaceBefore=4, spaceAfter=8))
    if not story:
        story.append(Spacer(1, 1))
    doc.build(story)


def safe_file_name(name: str, output_format: str, fallback: str = "dokument") -> str:
    base = os.path.splitext(os.path.basename((name or "").strip()))[0]
    base = re.sub(r"[^\w\-. ]+", "_", base, flags=re.UNICODE).strip(" ._")[:80] or fallback
    return f"{base}.{output_format}"


def build_file(content: str, output_format: str, out_path: str, title: str, font_path: str, bold_font_path: str, csv_delimiter: str) -> dict:
    """Inhalt in das gewünschte Format rendern. Gibt Kennzahlen für den Bericht zurück."""
    nodes = parse_markdown(content)
    tables = extract_tables(nodes)
    info = {"blocks": len(nodes), "tables": len(tables)}

    if output_format in ("xlsx", "csv"):
        table_rows: list[tuple[str, list[list[str]]]] = []
        last_heading = ""
        for node in nodes:
            if node.kind == "heading":
                last_heading = plain_text(node.text)
            elif node.kind == "table":
                table_rows.append((last_heading, node.rows))
                last_heading = ""
        if not table_rows:
            delimited = parse_delimited(content)
            if delimited:
                table_rows = [("", delimited)]
        if not table_rows:
            raise ValueError(
                "Für XLSX oder CSV wird eine Tabelle benötigt. Bitte den Inhalt als Markdown-Tabelle "
                "(| Spalte | Spalte |) oder als CSV-Zeilen übergeben."
            )
        if output_format == "xlsx":
            render_xlsx(table_rows, out_path)
        else:
            with open(out_path, "w", encoding="utf-8-sig", newline="") as fh:
                fh.write(render_csv(table_rows[0][1], csv_delimiter))
        info["tables"] = len(table_rows)
        return info

    if output_format == "docx":
        render_docx(nodes, out_path, title)
    elif output_format == "pdf":
        render_pdf(nodes, out_path, title, font_path, bold_font_path)
    elif output_format == "md":
        with open(out_path, "w", encoding="utf-8") as fh:
            fh.write(render_markdown(nodes, content))
    else:
        with open(out_path, "w", encoding="utf-8") as fh:
            fh.write(render_text(nodes))
    return info


# ---------------------------------------------------------------------------
# Open-WebUI-Tool
# ---------------------------------------------------------------------------


class Tools:
    class Valves(BaseModel):
        MAX_CONTENT_CHARS: int = Field(default=200_000, description="Obergrenze für den Inhalt in Zeichen.")
        FONT_PATH: str = Field(default=DEFAULT_FONT_PATH, description="TTF für PDF. Fehlt sie, Helvetica.")
        BOLD_FONT_PATH: str = Field(default=DEFAULT_BOLD_FONT_PATH, description="Fette TTF für PDF-Überschriften.")
        CSV_DELIMITER: str = Field(default=";", description="Trennzeichen für CSV. Semikolon passt zu deutschem Excel.")

    class UserValves(BaseModel):
        DEFAULT_OUTPUT_FORMAT: str = Field(default="docx", description="Format, wenn der Nutzer keines nennt.")

    def __init__(self) -> None:
        self.valves = self.Valves()

    async def create_document(
        self,
        content: str,
        file_name: str = "",
        output_format: str = "",
        title: str = "",
        __user__: Optional[dict] = None,
        __event_emitter__: Any = None,
    ) -> str:
        """
        Erzeugt aus einem vollständigen Markdown-Text eine Datei zum Download: Word (docx), PDF (pdf), Excel (xlsx), Markdown (md), Text (txt) oder CSV (csv). Nutze dieses Tool immer, wenn der Nutzer eine Datei, ein Dokument, eine Tabelle oder einen Export haben möchte. Schreibe den kompletten Inhalt selbst und übergib ihn in content. Für Excel oder CSV muss content eine Markdown-Tabelle enthalten.

        :param content: Der vollständige Inhalt der Datei als Markdown. Überschriften mit #, Listen mit -, Tabellen mit | Spalte | Spalte |, Fett mit **Text**.
        :param file_name: Gewünschter Dateiname ohne Endung, z. B. "Anschreiben Finanzamt". Leer = "dokument".
        :param output_format: docx, pdf, xlsx, md, txt oder csv. Leer = Standard des Nutzers.
        :param title: Dokumenttitel für die Metadaten. Leer = Dateiname.
        :return: Ergebnisbericht mit Download-Link. Gib dem Nutzer den Link unverändert weiter.
        """
        started = time.monotonic()
        user_valves = (__user__ or {}).get("valves")
        output_format = (output_format or "").strip().lower().lstrip(".") or getattr(
            user_valves, "DEFAULT_OUTPUT_FORMAT", "docx"
        )
        aliases = {"word": "docx", "excel": "xlsx", "markdown": "md", "text": "txt"}
        output_format = aliases.get(output_format, output_format)
        if output_format not in SUPPORTED_OUTPUT:
            return self._fail(f"Format '{output_format}' wird nicht unterstützt. Erlaubt: {', '.join(sorted(SUPPORTED_OUTPUT))}.")
        if not __user__ or not __user__.get("id"):
            return self._fail("Kein Nutzerkontext vorhanden. Abgebrochen.")
        if not content or not content.strip():
            return self._fail("Kein Inhalt übergeben. Bitte den vollständigen Text in content übergeben.")
        if len(content) > self.valves.MAX_CONTENT_CHARS:
            return self._fail(f"Inhalt zu groß ({len(content)} Zeichen, erlaubt {self.valves.MAX_CONTENT_CHARS}).")

        out_name = safe_file_name(file_name, output_format)
        title = (title or "").strip() or os.path.splitext(out_name)[0]
        await self._status(__event_emitter__, f"Erzeuge {out_name} …")

        try:
            record = await self._build_and_store(content, output_format, out_name, title, __user__["id"])
        except ValueError as exc:
            return self._fail(str(exc))
        except Exception as exc:
            log.warning("create_document: render/store failed (%s)", type(exc).__name__)
            return self._fail(f"Datei konnte nicht erzeugt werden: {type(exc).__name__}")

        url = f"/api/v1/files/{record['id']}/content"
        if __event_emitter__:
            await __event_emitter__(
                {
                    "type": "files",
                    "data": {
                        "files": [
                            {
                                "type": "file",
                                "id": record["id"],
                                "url": url,
                                "name": out_name,
                                "size": record["size"],
                                "content_type": record["content_type"],
                            }
                        ]
                    },
                }
            )
        await self._status(__event_emitter__, "Datei fertig.", done=True)
        return json.dumps(
            {
                "status": "success",
                "message": f"Datei erstellt. Download: [{out_name}]({url})",
                "download_url": url,
                "file_name": out_name,
                **record["info"],
                "seconds": round(time.monotonic() - started, 1),
                "instruction_for_assistant": "Zeige dem Nutzer den Download-Link als Markdown-Link. Wiederhole den Inhalt nicht.",
            },
            ensure_ascii=False,
        )

    # ----------------------------------------------------------------- helpers

    @staticmethod
    def _fail(message: str) -> str:
        return json.dumps({"status": "error", "message": message}, ensure_ascii=False)

    @staticmethod
    async def _status(emitter: Any, description: str, done: bool = False) -> None:
        if not emitter:
            return
        try:
            await emitter({"type": "status", "data": {"description": description, "done": done}})
        except Exception:
            pass

    async def _build_and_store(self, content: str, output_format: str, out_name: str, title: str, user_id: str) -> dict:
        from open_webui.models.files import FileForm, Files
        from open_webui.storage.provider import Storage

        with tempfile.TemporaryDirectory() as tmp:
            tmp_path = os.path.join(tmp, out_name)
            info = await asyncio.to_thread(
                build_file,
                content,
                output_format,
                tmp_path,
                title,
                self.valves.FONT_PATH,
                self.valves.BOLD_FONT_PATH,
                self.valves.CSV_DELIMITER,
            )
            with open(tmp_path, "rb") as fh:
                payload = fh.read()

        file_id = str(uuid.uuid4())
        content_type = CONTENT_TYPES[output_format]
        tags = {"OpenWebUI-User-Id": user_id, "OpenWebUI-File-Id": file_id}
        contents, stored_path = await asyncio.to_thread(
            Storage.upload_file, io.BytesIO(payload), f"{file_id}_{out_name}", tags
        )
        record = await Files.insert_new_file(
            user_id,
            FileForm(
                id=file_id,
                filename=out_name,
                path=stored_path,
                data={},
                meta={
                    "name": out_name,
                    "content_type": content_type,
                    "size": len(contents),
                    "data": {"generated_by": "create_document"},
                },
            ),
        )
        if record is None:
            raise RuntimeError("Dateieintrag konnte nicht angelegt werden.")
        return {"id": file_id, "size": len(contents), "content_type": content_type, "info": info}
