"""Tests für den Kern von create_document, ohne Open WebUI. Alle Inhalte sind erfunden.

Ausführen:
    python create_document/tests/test_core.py
"""

from __future__ import annotations

import os
import sys
import tempfile

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))

import create_document_tool as t  # noqa: E402

SAMPLE_MD = """# Prüfbericht Musterfirma Beispiel GmbH

Dieser Bericht ist **frei erfunden** und dient nur dem *Test*. Code wie `x = 1` bleibt Code.

## Feststellungen

- Erste Feststellung mit Umlauten: Größe, Straße
- Zweite Feststellung
  mit Fortsetzungszeile

1. Schritt eins
2. Schritt zwei

## Beträge

| Position | Betrag | Anzahl |
|---|---:|---|
| Honorar | 1.234,56 € | 2 |
| Auslagen | 99.5 | 3 |

---

```
rohdaten = "unverändert"
```

Schlussabsatz.
"""


def test_parse_markdown_structure():
    nodes = t.parse_markdown(SAMPLE_MD)
    kinds = [n.kind for n in nodes]
    assert kinds == [
        "heading", "paragraph", "heading", "bullets", "numbers", "heading", "table", "rule", "code", "paragraph",
    ], kinds
    assert nodes[0].level == 1 and nodes[2].level == 2
    assert nodes[3].items == ["Erste Feststellung mit Umlauten: Größe, Straße", "Zweite Feststellung mit Fortsetzungszeile"]
    assert nodes[4].items == ["Schritt eins", "Schritt zwei"]
    assert nodes[6].rows == [["Position", "Betrag", "Anzahl"], ["Honorar", "1.234,56 €", "2"], ["Auslagen", "99.5", "3"]]
    assert nodes[8].text == 'rohdaten = "unverändert"'


def test_inline_runs():
    runs = t.parse_inline("a **fett** b *kursiv* c `code` d")
    assert [(r.text, r.bold, r.italic, r.code) for r in runs] == [
        ("a ", False, False, False),
        ("fett", True, False, False),
        (" b ", False, False, False),
        ("kursiv", False, True, False),
        (" c ", False, False, False),
        ("code", False, False, True),
        (" d", False, False, False),
    ]
    assert t.plain_text("**x** und *y*") == "x und y"
    assert t.plain_text("2 * 3 * 4") == "2 * 3 * 4", "Sternchen mit Leerzeichen sind kein Markdown"


def test_cell_values_keep_german_numbers_as_text():
    assert t.cell_value("42") == 42
    assert t.cell_value("99.5") == 99.5
    assert t.cell_value("-3") == -3
    assert t.cell_value("1.234,56 €") == "1.234,56 €"
    assert t.cell_value("**Summe**") == "Summe"


def test_parse_delimited():
    assert t.parse_delimited("a;b;c\n1;2;3\n4;5;6") == [["a", "b", "c"], ["1", "2", "3"], ["4", "5", "6"]]
    assert t.parse_delimited("nur eine Zeile") is None
    assert t.parse_delimited("a,b\n1,2,3") is None, "ungleiche Spaltenzahl ist keine Tabelle"


def test_safe_file_name():
    assert t.safe_file_name("Anschreiben Finanzamt", "docx") == "Anschreiben Finanzamt.docx"
    assert t.safe_file_name("../../etc/passwd", "pdf") == "passwd.pdf"
    assert t.safe_file_name("", "xlsx") == "dokument.xlsx"
    assert t.safe_file_name("bericht.docx", "pdf") == "bericht.pdf"


def _build(fmt: str, content: str = SAMPLE_MD) -> tuple[str, dict]:
    tmp = tempfile.mkdtemp()
    path = os.path.join(tmp, f"out.{fmt}")
    info = t.build_file(content, fmt, path, "Titel", t.DEFAULT_FONT_PATH, t.DEFAULT_BOLD_FONT_PATH, ";")
    return path, info


def test_docx_output():
    import docx

    path, info = _build("docx")
    document = docx.Document(path)
    texts = [p.text for p in document.paragraphs if p.text]
    assert texts[0] == "Prüfbericht Musterfirma Beispiel GmbH"
    assert document.paragraphs[0].style.name.startswith("Heading")
    bold_runs = [r.text for p in document.paragraphs for r in p.runs if r.bold]
    assert "frei erfunden" in bold_runs
    assert any(p.style.name == "List Bullet" for p in document.paragraphs)
    assert any(p.style.name == "List Number" for p in document.paragraphs)
    assert len(document.tables) == 1
    assert document.tables[0].cell(1, 1).text == "1.234,56 €"
    assert info["tables"] == 1 and info["blocks"] == 10


def test_pdf_output_with_and_without_ttf():
    from pypdf import PdfReader

    path, _ = _build("pdf")
    text = "\n".join(page.extract_text() for page in PdfReader(path).pages)
    for needle in ("Prüfbericht", "Größe, Straße", "1.234,56 €", "Schritt zwei", "rohdaten", "Schlussabsatz."):
        assert needle in text, needle

    tmp = tempfile.mkdtemp()
    fallback = os.path.join(tmp, "fallback.pdf")
    t.build_file(SAMPLE_MD, "pdf", fallback, "T", "/nonexistent.ttf", "/nonexistent.ttf", ";")
    assert "Straße" in "\n".join(p.extract_text() for p in PdfReader(fallback).pages)


def test_xlsx_from_markdown_tables_and_from_csv_text():
    from openpyxl import load_workbook

    path, info = _build("xlsx")
    workbook = load_workbook(path)
    assert workbook.sheetnames == ["Beträge"]
    sheet = workbook["Beträge"]
    assert [c.value for c in sheet[1]] == ["Position", "Betrag", "Anzahl"]
    assert sheet["B2"].value == "1.234,56 €"
    assert sheet["B3"].value == 99.5 and sheet["C3"].value == 3
    assert sheet["A1"].font.bold
    assert info["tables"] == 1

    path, info = _build("xlsx", "Name;Wert\nA;1\nB;2")
    sheet = load_workbook(path).active
    assert sheet["A2"].value == "A" and sheet["B3"].value == 2

    try:
        _build("xlsx", "Nur ein Absatz ohne Tabelle.")
    except ValueError as exc:
        assert "Tabelle" in str(exc)
    else:
        raise AssertionError("expected ValueError")


def test_csv_output_uses_semicolon_and_bom():
    path, _ = _build("csv")
    with open(path, "rb") as fh:
        raw = fh.read()
    assert raw.startswith(b"\xef\xbb\xbf")
    assert raw.decode("utf-8-sig").splitlines()[0] == "Position;Betrag;Anzahl"


def test_txt_and_md_output():
    path, _ = _build("txt")
    text = open(path, encoding="utf-8").read()
    assert "Prüfbericht Musterfirma Beispiel GmbH\n====" in text
    assert "- Erste Feststellung" in text and "1. Schritt eins" in text
    assert "**" not in text
    path, _ = _build("md")
    assert open(path, encoding="utf-8").read().startswith("# Prüfbericht")


def test_empty_markdown_renders_valid_files():
    for fmt in ("docx", "pdf", "txt"):
        path, info = _build(fmt, "   ")
        assert os.path.getsize(path) > 0
        assert info["blocks"] == 0


if __name__ == "__main__":
    failures = 0
    for name, fn in sorted(globals().items()):
        if name.startswith("test_") and callable(fn):
            try:
                fn()
                print(f"PASS {name}")
            except Exception as exc:  # noqa: BLE001
                failures += 1
                print(f"FAIL {name}: {type(exc).__name__}: {exc}")
    sys.exit(1 if failures else 0)
